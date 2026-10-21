#!/usr/bin/env python3
"""Contrôle du classeur Excel : compare ses indicateurs (recalculés par Excel) aux valeurs du dashboard HTML.

Les valeurs « HTML » sont lues dans le dashboard (pages embarquées : tableaux, KPI, variables JavaScript) ;
les valeurs « Excel » sont lues dans une copie du classeur recalculée par Excel (cache des formules).
Produit un rapport d'écarts (Markdown + CSV) et signale toute cellule en erreur (#REF!, #VALUE!, #NAME?, #N/A, #DIV/0!).

Usage : python controle_dashboard_xlsx.py <dossier_resultats> <classeur_recalcule.xlsx> [--rapport rapport.md]
"""
from __future__ import annotations

import argparse
import csv
import glob
import json
import os
import re
import sys

import openpyxl
from bs4 import BeautifulSoup

DEC = json.JSONDecoder()
ERRORS = ("#REF!", "#VALUE!", "#NAME?", "#N/A", "#DIV/0!", "#NUM!", "#NULL!")
VIEWS = {"Référence": "Référence", "H1a": "H1a", "H1b": "H1b", "H1c": "H1c", "H2": "H2 (lissage)", "H3b": "H3b (lissage)"}


def num(text) -> float | None:
    """Extrait le 1er nombre d'un texte français (« 1 215 séances · 1 794 h » → 1215)."""
    if text is None:
        return None
    if isinstance(text, (int, float)):
        return float(text)
    t = re.sub(r"[\s  ]", "", str(text)).replace(",", ".")
    m = re.search(r"-?\d+(?:\.\d+)?", t)
    return float(m.group()) if m else None


def numbers(text: str) -> list[float]:
    t = re.sub(r"(?<=\d)[\s  ](?=\d{3})", "", str(text)).replace(",", ".")
    return [float(x) for x in re.findall(r"-?\d+(?:\.\d+)?", t)]


def load_pages(folder: str) -> dict:
    path = glob.glob(os.path.join(folder, "08_dashboard_complet_*.html"))[0]
    html = open(path, encoding="utf-8").read()
    raw = re.search(r'id="embedded-pages">(.*?)</script>', html, re.S).group(1).replace("<\\/", "</")
    return json.loads(raw)


def js_var(page: str, name: str):
    i = page.find(f"const {name} = ")
    return DEC.raw_decode(page[i + len(f"const {name} = "):].lstrip())[0] if i >= 0 else None


def html_tables(page: str) -> list[list[list[str]]]:
    soup = BeautifulSoup(re.sub(r"<script[^>]*>.*?</script>|<style.*?</style>", "", page, flags=re.S), "html.parser")
    out = []
    for t in soup.find_all("table"):
        rows = [[c.get_text(" ", strip=True) for c in tr.find_all(["td", "th"])] for tr in t.find_all("tr")]
        out.append(rows)
    return out


def sheet_values(wb, name: str, max_rows: int | None = None) -> dict:
    ws = wb[name]
    vals = {}
    for row in ws.iter_rows(min_row=1, max_row=max_rows):
        for c in row:
            if c.value is not None:
                vals[(c.row, c.column)] = c.value
    return vals


class Report:
    def __init__(self):
        self.rows = []

    def add(self, tab: str, item: str, html, excel, tol: float = 0.51):
        h = num(html) if not isinstance(html, float) else html
        e = num(excel)
        if h is None or e is None:
            status = "non comparable"
            diff = ""
        else:
            diff = e - h
            status = "OK" if abs(diff) <= tol else "ÉCART"
        self.rows.append({"onglet": tab, "élément": item, "HTML": html, "Excel": excel, "écart": diff, "statut": status})


def find_row(vals: dict, col: int, text: str, start: int = 1) -> int | None:
    for (r, c), v in sorted(vals.items()):
        if c == col and r >= start and str(v).strip() == text:
            return r
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("dossier")
    ap.add_argument("classeur")
    ap.add_argument("--rapport", default=None)
    args = ap.parse_args()
    folder = os.path.abspath(args.dossier)
    pages = load_pages(folder)
    wb = openpyxl.load_workbook(args.classeur, data_only=True, read_only=False)
    rep = Report()

    # --- 0. erreurs de formule dans les feuilles de restitution, Listes, Calcul (les tableaux de données sont contrôlés par Excel/COM) ---
    errors = []
    for ws in wb.worksheets:
        if ws.title.startswith("Données_"):
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in ERRORS:
                    errors.append((ws.title, c.coordinate, c.value))

    # --- 1. Conclusion LET ---
    ws = "Conclusion LET"
    v = sheet_values(wb, ws)
    t = html_tables(pages["conclusion"])
    for hrow in t[1][2:] if len(t[1]) > 2 else t[1][1:]:
        key = hrow[0]
        r = find_row(v, 2, key, 8)
        if r is None:
            continue
        # colonnes HTML : "n · h h" ; Excel : séances (col impaire) et heures
        names = ["À reporter", "En conflit avant", "Relogées", "Placées connu", "Placées estimé", "Restent connu", "Restent incertain", "Non placées sans effectif", "Non placées autre"]
        for j, name in enumerate(names):
            nums = numbers(hrow[1 + j])
            if len(nums) < 2:
                continue
            rep.add(ws, f"{key} · {name} (séances)", nums[0], v.get((r, 3 + 2 * j)))
            rep.add(ws, f"{key} · {name} (heures)", nums[1], v.get((r, 4 + 2 * j)), tol=1.0)

    # --- 2. Synthèse V2 : tableau des hypothèses et tableau de réaffectation ---
    ws = "Synthèse V2"
    v = sheet_values(wb, ws)
    t = html_tables(pages["synthese_v2"])
    for hrow in t[0][1:]:
        key = hrow[0].split(" : ")[0].strip()
        r = find_row(v, 2, hrow[0].strip().rstrip(), 1)
        if r is None:  # le libellé Excel est « H1a : description »
            for (rr, cc), val in v.items():
                if cc == 2 and str(val).startswith(key + " :"):
                    r = rr
                    break
        if r is None:
            continue
        col = 3
        for j, label in enumerate(["Locaux d'accueil", "À reporter", "dont examens", "Reportées connu", "Reportées estimé", "dont chevauchement", "Chevauchements", "Sans effectif", "Sans solution", "Non couvertes", "Besoin non résolu"]):
            nums = numbers(hrow[1 + j])
            if not nums:
                continue
            rep.add(ws, f"{key} · {label} (n)", nums[0], v.get((r, col)))
            col += 1
            if len(nums) > 1 and label not in ("Locaux d'accueil", "Chevauchements"):
                rep.add(ws, f"{key} · {label} (h)", nums[1], v.get((r, col)), tol=1.0)
                col += 1
    for hrow in t[1][1:]:
        sit = hrow[0].strip()
        r = find_row(v, 2, sit, 10)
        if r is None:
            continue
        rep.add(ws, f"{sit} · séances réaffectées", num(hrow[1]), v.get((r, 3)))
        rep.add(ws, f"{sit} · heures réaffectées", num(hrow[2]), v.get((r, 4)), tol=1.0)
        rep.add(ws, f"{sit} · heures d'amphi libérées (net)", num(hrow[3]), v.get((r, 5)), tol=1.0)
        occ, fill = numbers(hrow[4]), numbers(hrow[5])
        for (h, col, label) in ((occ[0], 6, "occupation avant"), (occ[1], 7, "occupation après"), (fill[0], 8, "remplissage avant"), (fill[1], 9, "remplissage après")):
            e = v.get((r, col))
            rep.add(ws, f"{sit} · {label} (%)", h, None if e is None else e * 100, tol=0.06)
        rep.add(ws, f"{sit} · effectif > capacité", num(hrow[6]), v.get((r, 10)))

    # --- 3. Usages sans effectif ---
    ws = "Usages sans effectif"
    v = sheet_values(wb, ws)
    soup = BeautifulSoup(pages["sans_effectif"], "html.parser")
    kp = {k.find("span").get_text(strip=True): k.find("strong").get_text(strip=True) for k in soup.select(".kpi")}
    for label, cell in (("Séances sans effectif", (5, 2)), ("Heures à reloger (8h-18h)", (5, 4)), ("Heures, durée complète", (5, 6))):
        rep.add(ws, label, kp.get(label), v.get(cell), tol=1.0)
    rep.add(ws, "Part des séances à reporter (%)", kp.get("Part des séances à reporter"), (v.get((5, 8)) or 0) * 100, tol=0.06)
    for hrow in html_tables(pages["sans_effectif"])[0][1:]:
        r = find_row(v, 2, hrow[0], 8)
        if r:
            rep.add(ws, f"{hrow[0]} · séances", num(hrow[1]), v.get((r, 3)))
            rep.add(ws, f"{hrow[0]} · heures 8h-18h", num(hrow[2]), v.get((r, 4)), tol=0.11)

    # --- 4. À vérifier ---
    ws = "À vérifier"
    v = sheet_values(wb, ws)
    t = html_tables(pages["verification"])
    for hrow in t[0][1:]:
        r = find_row(v, 2, hrow[0], 1)
        if r:
            rep.add(ws, f"Priorité {hrow[0]} · séances", num(hrow[1]), v.get((r, 3)))
            rep.add(ws, f"Priorité {hrow[0]} · heures", num(hrow[2]), v.get((r, 4)), tol=0.11)
    for hrow in t[2][1:]:
        r = find_row(v, 2, hrow[0], 1)
        if r:
            rep.add(ws, f"Sans salle · {hrow[0]} · séries", num(hrow[1]), v.get((r, 3)))
            rep.add(ws, f"Sans salle · {hrow[0]} · séances", num(hrow[2]), v.get((r, 4)))
            rep.add(ws, f"Sans salle · {hrow[0]} · heures", num(hrow[3]), v.get((r, 5)), tol=1.0)

    # --- 5. Vues de cas ---
    for scen, sheet in VIEWS.items():
        v = sheet_values(wb, sheet, max_rows=180)
        page = pages["reference" if scen == "Référence" else scen]
        raw = js_var(page, "rawData")
        conflicts = js_var(page, "rawConflicts") or []
        rooms = len(raw)
        tx = [r["taux"] for r in raw if r["taux"] is not None]
        kp = {
            "Salles suivies": rooms, "Taux moyen (%)": sum(tx) / len(tx), "Sous-utilisées (<10 %)": sum(1 for x in tx if x < 10),
            "Fortement sollicitées (≥30 %)": sum(1 for x in tx if x >= 30), "Heures-salles occupées": sum(r["heures"] for r in raw), "Séances reconstituées": sum(r["nbCours"] for r in raw),
            "Chevauchements détectés": len(conflicts) if scen != "Référence" else sum(1 for _ in conflicts) or None,
        }
        cells = {"Salles suivies": (8, 2), "Taux moyen (%)": (8, 4), "Sous-utilisées (<10 %)": (8, 6), "Fortement sollicitées (≥30 %)": (8, 8),
                 "Heures-salles occupées": (8, 10), "Séances reconstituées": (8, 12), "Chevauchements détectés": (8, 14)}
        for label, cell in cells.items():
            e = v.get(cell)
            if label == "Taux moyen (%)" and e is not None:
                e = e * 100
            tol = 0.06 if "%" in label else (1.0 if "Heures" in label else 0.51)
            if label == "Chevauchements détectés" and scen == "Référence":
                continue  # la référence n'embarque pas de liste de chevauchements dans le HTML (compte repris du CSV 28)
            rep.add(sheet, label, kp[label], e, tol)
        # occupation hebdomadaire : taux par semaine (HTML : weeklySalleData + jours ouvrés par semaine)
        weekly_salle = js_var(page, "weeklySalleData") or {}
        wk_days = js_var(page, "weeklyData")["TOUS"]
        hours_by_week, days_by_week = {}, {w["label"]: w["nbJoursOuvresSemaine"] for w in wk_days}
        for key, lst in weekly_salle.items():
            for w in lst:
                hours_by_week[w["label"]] = hours_by_week.get(w["label"], 0) + w["heures"]
        r0 = next((r for (r, c), x in sorted(v.items()) if c == 2 and x == "Semaine" and v.get((r, 3)) == "Clé"), None)
        if r0:
            diffs = []
            for k in range(1, 60):
                lab = v.get((r0 + k, 2))
                if lab is None or lab not in hours_by_week or not (v.get((r0 + k, 3)) or 0) > 100000:
                    continue
                expected = 100 * hours_by_week[lab] / max(1, rooms * days_by_week[lab] * 10)
                got = (v.get((r0 + k, 6)) or 0) * 100
                diffs.append(abs(got - expected))
            if diffs:
                rep.add(sheet, f"Occupation hebdomadaire : écart maximal sur {len(diffs)} semaines (points de %)", 0.0, max(diffs), tol=0.1)

    # --- 6. Amphis (7 cas × amphis) ---
    v = sheet_values(wb, "Amphis", max_rows=60)
    payload = json.loads(re.search(r'id="payload">(.*?)</script>', pages["amphis"], re.S).group(1).replace("<\\/", "</"))
    r_head = find_row(v, 3, "Amphi", 5)
    cases = payload["scenarios"]
    diffs = []
    for row in payload["rows"]:
        r = None
        for k in range(1, 60):
            if str(v.get((r_head + k, 4))) == str(row["code"]):
                r = r_head + k
                break
        if r is None:
            continue
        for j, case in enumerate(cases):
            exp = row["P"]["AN"][f"r_{case}"]
            got = (v.get((r, 10 + 2 * j)) or 0) * 100
            diffs.append(abs(exp - got))
    rep.add("Amphis", f"Taux de {len(payload['rows'])} amphis × {len(cases)} cas : écart maximal (points de %)", 0.0, max(diffs) if diffs else None, tol=0.011)

    # --- 7. Réaffectation globale : KPI de la situation Référence (valeur affichée à l'ouverture) ---
    v = sheet_values(wb, "Réaffectation globale", max_rows=15)
    sec = re.search(r'<div class="rz" data-h="Référence">(.*?)</div>\s*<div class="rz"', pages["reallocation"], re.S)
    soup = BeautifulSoup(sec.group(1) if sec else pages["reallocation"], "html.parser")
    kp = {k.find("span").get_text(strip=True): k.find("strong").get_text(strip=True) for k in soup.select(".kpi")}
    mapping = [("Cours/TD SSH Pau analysés", (8, 2), 1), ("Séances réaffectées", (8, 4), 1), ("Heures d'amphi libérées (net)", (8, 10), 1)]
    for label, cell, tol in mapping:
        rep.add("Réaffectation globale", f"Référence · {label}", kp.get(label), v.get(cell), tol=tol)

    # --- rapport ---
    ok = sum(1 for r in rep.rows if r["statut"] == "OK")
    bad = [r for r in rep.rows if r["statut"] == "ÉCART"]
    other = [r for r in rep.rows if r["statut"] == "non comparable"]
    lines = ["# Rapport de contrôle du classeur Excel", "",
             f"- Éléments comparés : {len(rep.rows)} (OK : {ok} ; écarts : {len(bad)} ; non comparables : {len(other)})",
             f"- Cellules en erreur (#REF!, #VALUE!, #NAME?, #N/A, #DIV/0! — feuilles de restitution, Listes, Calcul) : {len(errors)}", ""]
    if errors:
        lines += ["## Erreurs de formule", ""] + [f"- {s}!{c} : {e}" for s, c, e in errors[:50]] + [""]
    lines += ["## Écarts", ""]
    lines += ["Aucun écart au-delà de la tolérance d'arrondi." if not bad else "| Onglet | Élément | HTML | Excel | Écart |", "|---|---|---|---|---|"] if not bad else ["| Onglet | Élément | HTML | Excel | Écart |", "|---|---|---|---|---|"]
    lines += [f"| {r['onglet']} | {r['élément']} | {r['HTML']} | {r['Excel']} | {r['écart']:+.3f} |" for r in bad]
    if other:
        lines += ["", "## Éléments non comparables (valeur absente d'un côté)", ""] + [f"- {r['onglet']} · {r['élément']} : HTML={r['HTML']} / Excel={r['Excel']}" for r in other[:40]]
    report_path = args.rapport or os.path.join(folder, "rapport_controle_dashboard_xlsx.md")
    open(report_path, "w", encoding="utf-8").write("\n".join(lines) + "\n")
    with open(os.path.splitext(report_path)[0] + ".csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rep.rows[0]), delimiter=";")
        w.writeheader()
        w.writerows(rep.rows)
    print("\n".join(lines[:6]))
    print(f"Rapport : {report_path}")


if __name__ == "__main__":
    main()
